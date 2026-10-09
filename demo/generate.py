"""Generate the fictional Northpeak demo world (BUILD.md §17.1). Seeded and deterministic.

Writes:
  demo/docs/                  the knowledge documents (DOCX, Markdown, HTML, PDF), incl. Handbook v2
  demo/data/traffic.jsonl     ~2,400 prompts over 7 simulated days from 25 users, with ground truth
  demo/data/labelled_tickets.csv   300 labelled tickets for the live intake (S6)
  demo/data/urgency_train.jsonl    labelled urgency texts for the seeded `dl` artifact
  demo/data/eval/             labelled evaluation sets (router, cache pairs, retrieval)

Run: make demo-data   (python demo/generate.py)
"""
from __future__ import annotations

import csv
import json
import random
import sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "backend"))

from app import demo_world as w  # noqa: E402

SEED = 20251020
DATA = HERE / "data"
DOCS = HERE / "docs"
EVAL = DATA / "eval"
WEEK_S = 7 * 24 * 3600


# ----------------------------------------------------------------------------- documents
def render_docx(spec: w.DocSpec, path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(spec.title, 0)
    for heading, paras in spec.sections[:1]:
        d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)
    if spec.table:
        heading, header, rows = spec.table
        d.add_heading(heading, 1)
        d.add_paragraph("Every employee receives the following paid vacation days per year, by country of "
                        "employment.")
        t = d.add_table(rows=1, cols=len(header))
        for i, h in enumerate(header):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = v
    for heading, paras in spec.sections[1:]:
        d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)
    d.core_properties.title = spec.title
    d.save(str(path))


def render_md(spec: w.DocSpec, path: Path) -> None:
    lines = [f"# {spec.title}", ""]
    for heading, paras in spec.sections:
        lines += [f"## {heading}", ""]
        for p in paras:
            lines += [p, ""]
    path.write_text("\n".join(lines))


def render_html(spec: w.DocSpec, path: Path) -> None:
    body = [f"<h1>{spec.title}</h1>"]
    for heading, paras in spec.sections:
        body.append(f"<h2>{heading}</h2>")
        body += [f"<p>{p}</p>" for p in paras]
    path.write_text("<!doctype html><html><head><meta charset='utf-8'><title>"
                    f"{spec.title}</title></head><body>{''.join(body)}</body></html>")


def render_pdf(spec: w.DocSpec, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(spec.title, styles["Title"])]
    for heading, paras in spec.sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story += [Paragraph(p, styles["BodyText"]), Spacer(1, 6)]
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=spec.title, invariant=1)
    doc.build(story)


RENDER = {"docx": render_docx, "md": render_md, "html": render_html, "pdf": render_pdf}


def write_documents() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    for spec in [*w.DOCUMENTS, w.HANDBOOK_V2]:
        RENDER[spec.fmt](spec, DOCS / spec.filename)
    manifest = [{"key": d.key, "collection": d.collection, "title": d.title, "file": d.filename}
                for d in w.DOCUMENTS]
    (DOCS / "manifest.json").write_text(json.dumps({"documents": manifest, "handbook_v2": w.HANDBOOK_V2.filename},
                                                   indent=1))


# ----------------------------------------------------------------------------- traffic
MIX = [("triage", 0.30), ("extraction", 0.15), ("urgency", 0.10), ("conversion", 0.05),
       ("knowledge", 0.20), ("open", 0.17), ("routing", 0.03)]
TASK_TYPE = {"triage": "classification", "extraction": "extraction", "urgency": "scoring",
             "conversion": "transform", "routing": "routing", "vacation": "lookup", "knowledge": None,
             "open": None}


def user_can_read(user: dict, collection_key: str) -> bool:
    col = next(c for c in w.COLLECTIONS if c["key"] == collection_key)
    return col["access"] == "org" or bool(set(col["group_ids"]) & set(user["groups"]))


def gen_traffic(rng: random.Random, n: int = 2400) -> list[dict]:
    users = w.ALL_USERS
    weights = [3.0 if u in w.PERSONAS else 1.0 for u in users]
    rows: list[dict] = []
    tasks = [t for t, _ in MIX]
    probs = [p for _, p in MIX]
    triage_cls_w = {"billing": 0.2, "shipping": 0.22, "returns": 0.18, "refund": 0.1, "technical": 0.18,
                    "account": 0.12}
    for i in range(n):
        task = rng.choices(tasks, probs)[0]
        user = rng.choices(users, weights)[0]
        collections: list[str] = []
        truth: str | None
        if task == "triage":
            cls = rng.choices(list(triage_cls_w), list(triage_cls_w.values()))[0]
            prompt, _ = w.make_triage_prompt(rng, cls)
            truth = cls
            assert w.classify_ticket(w.quoted(prompt)) == cls, prompt
        elif task == "extraction":
            prompt = w.make_extraction(rng)
            truth = w.extract_order(w.quoted(prompt))
        elif task == "urgency":
            prompt = w.make_urgency(rng)
            truth = str(w.urgency_score(w.quoted(prompt)))
        elif task == "conversion":
            prompt = w.make_conversion(rng)
            truth = w.convert_units(prompt)
        elif task == "routing":
            prompt = w.make_routing(rng)
            truth = w.route_queue(prompt)
        elif task == "knowledge":
            if rng.random() < 0.3:
                country = rng.choice(list(w.VACATION_V1))
                prompt = rng.choice(w.VACATION_Q).format(c=country)
                truth = w.vacation_answer(country, w.VACATION_V1[country])
                task = "vacation"
                collections = ["handbook"]
            else:
                readable = [q for q in w.KNOWLEDGE_QUESTIONS if user_can_read(user, q[1])]
                q, col, frag = rng.choice(readable)
                prompt, truth = q, frag
                collections = [c["key"] for c in w.COLLECTIONS if user_can_read(user, c["key"])]
        else:
            prompt = w.make_open_ended(rng)
            truth = None
        ts = int(i * WEEK_S / n + rng.randint(0, 120))
        rows.append({"i": i, "t": ts, "user_id": user["id"], "prompt": prompt, "task": task,
                     "task_type": TASK_TYPE[task], "truth": truth, "collections": collections})
    return rows


def gen_labelled_tickets(rng: random.Random, n: int = 300) -> list[tuple[str, str]]:
    out = []
    classes = w.TRIAGE_CLASSES
    for i in range(n):
        cls = classes[i % len(classes)]
        out.append((w.make_ticket(rng, cls), cls))
    rng.shuffle(out)
    return out


def gen_urgency_train(rng: random.Random, n: int = 1200) -> list[dict]:
    seen: set[str] = set()
    out = []
    while len(out) < n:
        t = w.make_urgency_text(rng)
        prompt = rng.choice(w.URGENCY_PROMPTS).format(t=t)
        if prompt in seen and rng.random() < 0.7:
            continue
        seen.add(prompt)
        out.append({"prompt": prompt, "label": str(w.urgency_score(t))})
    return out


# ----------------------------------------------------------------------------- evaluation sets
def gen_eval(rng: random.Random) -> dict[str, list[dict]]:
    must_match = []
    for _ in range(60):
        must_match.append({"prompt": w.make_extraction(rng), "pattern": "Order details extraction"})
    for _ in range(40):
        must_match.append({"prompt": w.make_conversion(rng), "pattern": "Unit conversion"})
    for _ in range(40):
        must_match.append({"prompt": w.make_urgency(rng), "pattern": "Urgency scoring"})

    must_not = []
    for _ in range(40):
        must_not.append({"prompt": w.make_open_ended(rng), "reason": "generation"})
    for q, _, _ in w.KNOWLEDGE_QUESTIONS:
        must_not.append({"prompt": q, "reason": "free-form knowledge question"})
    for p in ["Translate 'good morning' into French", "What's the weather like in Chamonix?",
              "Tell me a joke about tents", "Explain how down insulation works",
              "Write a haiku about hiking boots", "What should I pack for a 3-day hike?",
              "Compare the Ridgeline 2 with a hammock", "Who founded Northpeak?",
              "Is it safe to camp above the tree line in a storm?", "List five tips for winter camping"]:
        must_not.append({"prompt": p, "reason": "off-pattern"})

    cache_pairs = []
    for n1, n2 in [(5, 50), (12, 120), (3, 30), (10, 100), (1, 2), (8, 80), (20, 25)]:
        cache_pairs.append({"first": f"Convert {n1} km to miles", "second": f"Convert {n2} km to miles",
                            "should_hit": False, "reason": "slot mismatch"})
    cache_pairs += [
        {"first": "convert 5 USD to EUR", "second": "convert 50 USD to EUR", "should_hit": False,
         "reason": "slot mismatch"},
        {"first": "What's the status of today's shipments?", "second": "What's the status of today's shipments?",
         "should_hit": False, "reason": "time-sensitive"},
        {"first": "What is the latest price of the Ridgeline 2?", "second": "What is the latest price of the Ridgeline 2?",
         "should_hit": False, "reason": "time-sensitive"},
        {"first": "Draft a friendly reply to a customer whose tent pole broke",
         "second": "Draft a friendly reply to a customer whose tent pole broke", "should_hit": True,
         "reason": "exact repeat"},
        {"first": "Draft a friendly reply to a customer whose tent pole broke",
         "second": "draft a friendly reply to a customer whose tent pole broke!", "should_hit": True,
         "reason": "normalized repeat"},
        {"first": "Write a product description for the backpack",
         "second": "Write a product description for the headlamp", "should_hit": False,
         "reason": "different product"},
    ]

    retrieval = []
    for q, col, frag in w.KNOWLEDGE_QUESTIONS:
        doc = next(d for d in w.DOCUMENTS if d.collection == col and frag.lower() in w.doc_plain_text(d).lower())
        retrieval.append({"question": q, "collection": col, "document": doc.title, "answer_fragment": frag})
    for c in ["France", "Germany", "Spain", "Sweden"]:
        retrieval.append({"question": w.VACATION_Q[0].format(c=c), "collection": "handbook",
                          "document": "Employee Handbook", "answer_fragment": str(w.VACATION_V1[c])})

    followups = [
        {"previous": "How many vacation days in France?", "prompt": "and for Spain?"},
        {"previous": "How many vacation days do employees in France get?", "prompt": "What about Germany?"},
        {"previous": "Convert 12 km to miles", "prompt": "same for 20?"},
        {"previous": "How long is the warranty on Northpeak gear?", "prompt": "and does it cover that?"},
        {"previous": "What is the daily meal allowance on business trips?", "prompt": "How about hotels?"},
        {"previous": "Classify this support ticket: 'my tent pole snapped'", "prompt": "it was the second one"},
    ]
    return {"router_must_match": must_match, "router_must_not_match": must_not, "cache_pairs": cache_pairs,
            "retrieval_qa": retrieval, "followups": followups}


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows))


def main() -> None:
    rng = random.Random(SEED)
    DATA.mkdir(parents=True, exist_ok=True)
    EVAL.mkdir(parents=True, exist_ok=True)
    write_documents()

    traffic = gen_traffic(rng)
    write_jsonl(DATA / "traffic.jsonl", traffic)

    tickets = gen_labelled_tickets(random.Random(SEED + 1))
    with (DATA / "labelled_tickets.csv").open("w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["ticket_text", "label"])
        wr.writerows(tickets)

    write_jsonl(DATA / "urgency_train.jsonl", gen_urgency_train(random.Random(SEED + 2)))
    for name, rows in gen_eval(random.Random(SEED + 3)).items():
        write_jsonl(EVAL / f"{name}.jsonl", rows)

    triage = [r for r in traffic if r["task"] == "triage"]
    kw = sum(w.keyword_rule(w.quoted(r["prompt"])) == r["truth"] for r in triage) / max(1, len(triage))
    mix = Counter(r["task"] for r in traffic)
    print(f"traffic: {len(traffic)} prompts from {len({r['user_id'] for r in traffic})} users; mix {dict(mix)}")
    print(f"triage keyword-rule agreement: {kw:.1%} (target 60-75 %)")
    print(f"documents: {sorted(p.name for p in DOCS.iterdir())}")


if __name__ == "__main__":
    main()
